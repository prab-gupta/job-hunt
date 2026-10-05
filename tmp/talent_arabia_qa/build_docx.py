from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

ROOT=Path('/Users/prab/Downloads/claud jobs')
SRC=ROOT/'cv/main_TalentArabia_SeniorAIEngineer.tex'
OUT=ROOT/'output/documents/Prabhavit_Gupta_TalentArabia_AI_Engineer_CV.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
D=Document()
S=D.sections[0]
S.page_width=Cm(21); S.page_height=Cm(29.7)
S.top_margin=Cm(1.6); S.bottom_margin=Cm(1.8); S.left_margin=Cm(1.9); S.right_margin=Cm(1.9)
S.footer_distance=Cm(0.75)
WIDTH=S.page_width-S.left_margin-S.right_margin
for name in ['Normal','Title','Heading 1','List Bullet','List Bullet 2']:
    st=D.styles[name]
    st.font.name='Liberation Serif'; st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(10)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'Liberation Serif')
    st._element.rPr.rFonts.set(qn('w:hAnsi'),'Liberation Serif')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        st._element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    for border in list(st._element.xpath('./w:pPr/w:pBdr')): border.getparent().remove(border)
    st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(2)
    st.paragraph_format.line_spacing=Pt(12)
D.styles['Title'].font.size=Pt(24); D.styles['Title'].font.bold=True
D.styles['Title'].paragraph_format.line_spacing=Pt(28)
D.styles['Heading 1'].font.size=Pt(14.4); D.styles['Heading 1'].font.bold=True
D.styles['Heading 1'].font.color.rgb=RGBColor.from_string('0000B3')
D.styles['Heading 1'].paragraph_format.space_before=Pt(9)
D.styles['Heading 1'].paragraph_format.space_after=Pt(5)
D.styles['Heading 1'].paragraph_format.line_spacing=Pt(17)
D.styles['Heading 1'].paragraph_format.keep_with_next=True
D.core_properties.title='Prabhavit Gupta CV for Talent Arabia AI Engineer'
D.core_properties.subject='AI solutions and automation engineering experience'
D.core_properties.author='Prabhavit Gupta'
D.core_properties.keywords='AI Agents, RAG, n8n, Eden AI, API integration, automation'


def group(s,i):
    assert s[i]=='{',(s,i)
    level=1; j=i+1
    while level:
        if s[j]=='{' and s[j-1]!='\\': level+=1
        if s[j]=='}' and s[j-1]!='\\': level-=1
        j+=1
    return s[i+1:j-1],j


def plain(s):
    s=s.replace(r'\&','&').replace('---','-').replace('--','-')
    return s


def add_link(p,text,url,bold=False,italic=False):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'),'Liberation Serif'); fonts.set(qn('w:hAnsi'),'Liberation Serif'); pr.append(fonts)
    col=OxmlElement('w:color'); col.set(qn('w:val'),'0000B3'); pr.append(col)
    if bold: pr.append(OxmlElement('w:b'))
    if italic: pr.append(OxmlElement('w:i'))
    r.append(pr); t=OxmlElement('w:t'); t.text=plain(text); r.append(t); h.append(r); p._p.append(h)


def formatted(p,text,bold=False,italic=False):
    i=0; buf=''
    def flush():
        nonlocal buf
        if buf:
            r=p.add_run(plain(buf)); r.bold=bold; r.italic=italic; buf=''
    while i<len(text):
        typ=next((x for x in ['\\textbf','\\textit','\\href'] if text.startswith(x+'{',i)),None)
        if not typ:
            buf+=text[i]; i+=1; continue
        flush(); content,j=group(text,i+len(typ))
        if typ=='\\href':
            label,j=group(text,j); add_link(p,label,content,bold,italic)
        else:
            formatted(p,content,bold or typ=='\\textbf',italic or typ=='\\textit')
        i=j
    flush()


def heading(text):
    p=D.add_paragraph(plain(text),'Heading 1')
    # Original CV section design uses a blue 0.5 pt rule.
    b=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
    for k,v in [('val','single'),('sz','4'),('space','1'),('color','0000B3')]: bottom.set(qn('w:'+k),v)
    b.append(bottom); p._p.get_or_add_pPr().append(b)


def bullet(text,depth):
    p=D.add_paragraph()
    f=p.paragraph_format
    f.left_indent=Pt(11 if depth==1 else 22)
    f.first_line_indent=Pt(-11 if depth==1 else -10)
    p.add_run('• ' if depth==1 else '- ')
    formatted(p,text)
    if text.startswith('\\textbf{Product') or text.startswith('\\textbf{Technical'): f.keep_with_next=True


def role(line):
    args=[]; pos=len('\\role')
    for _ in range(4):
        value,pos=group(line,pos); args.append(value)
    org,title,place,date=args
    # Role metadata is kept in the same two aligned columns as the CV template.
    table=D.add_table(rows=2,cols=2)
    table.autofit=False
    table.columns[0].width=Cm(6.6); table.columns[1].width=Cm(10.6)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge); b.set(qn('w:val'),'nil'); borders.append(b)
    pr.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        m=OxmlElement('w:'+edge); m.set(qn('w:w'),'0'); m.set(qn('w:type'),'dxa'); margin.append(m)
    pr.append(margin)
    for row in table.rows:
        for cell in row.cells:
            cp=cell._tc.get_or_add_tcPr()
            # Natural cell wrapping preserves long education and role titles.
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=True
    p=table.cell(0,0).paragraphs[0]; p.paragraph_format.space_before=Pt(3)
    formatted(p,org,bold=True)
    p=table.cell(0,1).paragraphs[0]; p.paragraph_format.space_before=Pt(3)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; formatted(p,title,italic=True)
    p=table.cell(1,0).paragraphs[0]; p.paragraph_format.space_after=Pt(1); formatted(p,place,italic=True)
    p=table.cell(1,1).paragraphs[0]; p.paragraph_format.space_after=Pt(1)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; formatted(p,date,italic=True)


# Float the original profile photo at the right of the title and contact area.
p=D.add_paragraph('Prabhavit Gupta','Title'); p.paragraph_format.space_after=Pt(8)
p.paragraph_format.right_indent=Cm(3.0)
r=p.add_run(); shape=r.add_picture(str(ROOT/'cv/photo.jpg'),width=Cm(2.3),height=Cm(2.3))
inline=shape._inline; extent=inline.find(qn('wp:extent')); graphic=inline.find(qn('a:graphic'))
anchor=OxmlElement('wp:anchor')
for k,v in [('distT','0'),('distB','0'),('distL','0'),('distR','0'),('simplePos','0'),('relativeHeight','0'),('behindDoc','0'),('locked','0'),('layoutInCell','1'),('allowOverlap','1')]: anchor.set(k,v)
simple=OxmlElement('wp:simplePos'); simple.set('x','0'); simple.set('y','0'); anchor.append(simple)
ph=OxmlElement('wp:positionH'); ph.set('relativeFrom','margin'); align=OxmlElement('wp:align'); align.text='right'; ph.append(align); anchor.append(ph)
pv=OxmlElement('wp:positionV'); pv.set('relativeFrom','margin'); off=OxmlElement('wp:posOffset'); off.text='0'; pv.append(off); anchor.append(pv)
anchor.append(extent)
eff=OxmlElement('wp:effectExtent'); [eff.set(k,'0') for k in ['l','t','r','b']]; anchor.append(eff)
anchor.append(OxmlElement('wp:wrapNone'))
pr=inline.find(qn('wp:docPr')); pr.set('descr','Prabhavit Gupta profile photograph'); anchor.append(pr)
fr=inline.find(qn('wp:cNvGraphicFramePr')); anchor.append(fr); anchor.append(graphic)
inline.getparent().replace(inline,anchor)
p=D.add_paragraph(); p.paragraph_format.right_indent=Cm(2.7); p.paragraph_format.space_after=Pt(28)
p.paragraph_format.line_spacing=Pt(11)
r=p.add_run('+39-351-610-7807   '); r.font.size=Pt(9)
add_link(p,'prabhavitg@gmail.com','mailto:prabhavitg@gmail.com')
p.add_run('   '); add_link(p,'LinkedIn','https://www.linkedin.com/in/prabhavit/')
p.add_run('   Venice, Italy')
for r in p.runs: r.font.size=Pt(9)

body=SRC.read_text().split('\\section{Career Summary}',1)[1]
body='\\section{Career Summary}'+body
level=0
for line in body.splitlines():
    line=line.strip()
    if line.startswith('\\section{'):
        heading(group(line,len('\\section'))[0])
    elif line.startswith('\\begin{itemize}'): level+=1
    elif line.startswith('\\end{itemize}'): level-=1
    elif line.startswith('\\item '): bullet(line[6:],level)
    elif line.startswith('\\role{'): role(line)
    elif line=='\\newpage':
        p=D.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=Pt(1)
        p.add_run().add_break(WD_BREAK.PAGE)

f=S.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.paragraph_format.line_spacing=Pt(11)
r=f.add_run('Page '); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('808080')
for text in ['PAGE',' of ','NUMPAGES']:
    if text==' of ':
        r=f.add_run(text); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('808080')
    else:
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),text)
        rr=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
        sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'18'); rp.append(sz)
        co=OxmlElement('w:color'); co.set(qn('w:val'),'808080'); rp.append(co)
        rr.append(rp); t=OxmlElement('w:t'); t.text='1' if text=='PAGE' else '2'; rr.append(t)
        fld.append(rr); f._p.append(fld)
D.save(OUT)
print(OUT)
