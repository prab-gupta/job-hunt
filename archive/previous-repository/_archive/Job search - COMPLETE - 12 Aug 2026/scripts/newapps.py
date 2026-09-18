import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.load_workbook('Role shortlist - 8 Aug 2026.xlsx')
ws = wb['Applications sent']

GREEN = PatternFill('solid', fgColor='D9EAD3')
AMBER = PatternFill('solid', fgColor='FCE5CD')
GREY  = PatternFill('solid', fgColor='EFEFEF')
FILL  = {'SUBMITTED': GREEN, 'NEEDS YOU': AMBER, 'HELD': AMBER, 'CLOSED': GREY, 'DROPPED': GREY}
THIN  = Side(style='thin', color='BBBBBB')
BOX   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NEW = [
 ['opus','Founding Engineer','Berlin, Germany (hybrid)','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to opus."',
  'join.com, ~2 min. Cover letter attached. Start 1 Sep 2026, comp 40000 EUR, onsite Yes. No years gate and German is only a Nice-to-Have. $6M from Project A, Speedinvest and Sequoia Scouts; AI agents for industrial supply chain. Letter states the gaps: no FastAPI/Pydantic/SQLAlchemy specifically, and agentic work via n8n rather than PydanticAI/Google ADK/LangGraph. Contact person Tarkan Akbogra.',
  'https://join.com/companies/opus-techai/15525100'],
 ['Avery','Founding Engineer (Core)','Amsterdam, Netherlands (hybrid)','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to Avery."',
  'join.com. Their advertised band is EUR 3,500-5,000/month (42-60k/yr) which is ABOVE your stated 30-40k, but I entered 40000 because your number is yours to set, not mine. Worth knowing you undershot their floor. NL work authorisation answered Yes per your EU long-term permit. Free-text stack question answered with the real stack plus the two gaps (agent work in n8n not a Python framework, no CS degree). Founder Alisher Jafarov. Real founding equity.',
  'https://join.com/companies/getaveryai/16508512-founding-engineer-core'],
 ['The Quality Group (ESN / More)','AI Engineer (gn)','Deutschland, remote - German residency required','SUBMITTED','12 Aug 2026',
  '"Vielen Dank fuer deine Bewerbung! / Thanks for your application!"',
  'BEST STACK MATCH OF THE WHOLE SWEEP. The posting names n8n, Claude Enterprise and ChatGPT Enterprise as the tools. No German language requirement anywhere. Cover letter written for it. Greenhouse autofilled the OLD Prabhavit_CV-.pdf again - removed and replaced. Used the free notes field to raise the German residency point head on and ask them to tell you if day-one residency is a hard gate. 1,200+ employees, sports nutrition.',
  'https://job-boards.greenhouse.io/thequalitygroupgmbh2/jobs/4931097101'],
 ['Reedsy','AI Engineer - Remote Europe','UK; France; NL; Poland; Portugal; Spain; Remote','SUBMITTED','12 Aug 2026',
  '"Your application was successfully submitted! We will share an update as soon as possible."',
  'HIGHEST PAY IN THE SWEEP: EUR 70,000-120,000 + 0-0.2% equity, published on the posting. No years bar. 5,000,000 authors. Building a narrative co-pilot. Italy is not in their country list but "Remote" is, so I applied and let them decide. Long free-text answer names four gaps outright: no LangChain/LlamaIndex, no GCP, no Elasticsearch, no fine-tuning, plus ~3 yrs and no CS degree.',
  'https://jobs.ashbyhq.com/reedsy/b043b550-02b7-416f-8969-5db8fd9dc8ee'],
 ['Spyrosoft','Forward Deployed Engineer (OpenAI)','Poland (PL), fully remote','SUBMITTED','12 Aug 2026',
  '"Thanks for applying - We have received your application and we will be reviewing it shortly."',
  'EXACT TARGET TITLE AT YOUR LEVEL. Posting says "Depending on your experience level (Regular or Senior)" so Regular is in scope. PLN 110-150/hour, suits the Partita IVA. Fluent English C1, no Polish. Nice-to-haves you lack: Azure OpenAI, LangChain/LangGraph, Docker/K8s. 1,800+ staff, EUR 104M turnover. Phone selector defaulted to +48 - switched to +39 first. Left the post-submit profile-creation consent unticked (that creates an account).',
  'https://careers.spyro-soft.com/jobs/8044521-forward-deployed-engineer-openai'],
 ['Hexa (Waniwani)','Forward Deployed Engineer','Remote-first, based in Europe; travel Paris/London/Madrid','SUBMITTED','12 Aug 2026',
  '"Thanks for your interest in Hexa, we\'ll take the time to properly review your profile."',
  'AI assistants distributing insurance and banking products. $8M seed led by Seedcamp. Explicitly NOT a hands-on coding role - 35% discovery, 35% architecture, 30% procurement and security; a freelancer network implements. No years bar, first hire in the function. Free-text answer states two gaps outright: no regulated-industry background and never driven a security questionnaire or procurement committee. Ashby combobox overwrote the location with "Itasy, Madagascar" mid-form; caught on screenshot and fixed.',
  'https://jobs.ashbyhq.com/hexa/9040c373-0434-4274-b40d-6a18c1f7f803'],
 ['Nuitée','GTM Engineer','Remote in EU, or hybrid in Barcelona','SUBMITTED','12 Aug 2026',
  '"Thank you for applying. Your application has been received."',
  'CLEANEST ELIGIBILITY LINE IN THE SWEEP - "Remote in EU" with no residency or nationality catch. Stripe for Travel. Profitable before a $48M Series A led by Accel, 10x since 2022; Hopper, Priceline, Google and Uber are customers. Reports to Gian, VP Growth, building the product-led GTM machine from a blank page. Their tools: PostHog, HubSpot, Clay, Looker Studio, BigQuery. Letter says you are at the bottom of their 3-5 yr band and that you have not used those specific tools.',
  'https://job-boards.greenhouse.io/nuitee/jobs/4901595101'],
 ['Pliant','Software Engineer - AI Experience','Remote (Germany)','SUBMITTED','12 Aug 2026',
  '"Your application was successfully submitted. We\'ll take a close look at your application and be in touch soon!"',
  'Corporate card / spend management fintech. No years number - the bar is "shipping AI-powered products into production". RAG, embeddings, tool calling, agentic loops, evals. Their primary stack is Java/Spring Boot and/or Python, so you are applying on the Python side, not your strongest. Salary field took your range verbatim. Posted within a day of applying.',
  'https://jobs.ashbyhq.com/pliant/b281fc8b-599a-4082-a48e-59084b3697c9'],
 ['HiPeople','Applied AI Engineer - Systems & Reliability','World / Remote, Berlin office optional','SUBMITTED','12 Aug 2026',
  '"Your application was successfully submitted. We\'ll contact you if there are next steps."',
  'Genuinely CV-only: name, email, resume, LinkedIn, nothing else. Location field literally says "World". No years bar, no degree gate; welcomes AI/ML, backend/systems or data science backgrounds with production experience. Six-step interview compressed into about a week.',
  'https://jobs.ashbyhq.com/hipeople-official/6c330d7b-7c6d-4993-8893-b58b5289d442'],
 ['Jamie','AI Engineer - top 5% salary band in Europe','Remote / Berlin Mitte several days a week','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to Jamie."',
  'AI meeting assistant, 11-50 people. "2+ years of development experience preferred" is the LOWEST bar found all sweep. Requirements are prompt engineering, ML/AI pipelines, MLOps and product mindset - not ML research. Updated 19 hours before applying. No cover-letter step on the form, so I used the city question to volunteer the permit and your willingness to relocate to Berlin. Contact Egor Spirin.',
  'https://join.com/companies/meetjamieai/16567852-ai-engineer-top-5-salary-band-in-europe'],
 ['otonomee','Full Stack AI Engineer (Ireland AND Portugal)','Remote in Ireland / Remote in Portugal','NEEDS YOU','12 Aug 2026','-',
  'TWO SEPARATE APPLICATIONS, both blocked on the same thing. The form makes Internet Provider, Speed Test Download Speed and Speed Test Upload Speed MANDATORY. Those are facts about your home connection and I will not invent them. Run a speedtest and fill three fields; everything else is CV-only. Worth doing: no years-of-experience bar at all, and the stack listed is your CV restated (OpenAI/Anthropic, RAG, vector DBs, agents, Next.js, webhooks, Zendesk/Intercom, Azure). First engineering hire, reports to CTO. Ireland: /job/2656102 - Portugal: /job/2656143',
  'https://otonomee.jobs.personio.com/job/2656102/apply?language=en'],
 ['Capco Poland','AI Agent Engineer with Copilot Studio and Power Platform','Poland, hybrid (1-2 days/month Krakow), B2B contract','NEEDS YOU','12 Aug 2026','-',
  'LIVE and it is literally your Microsoft certifications as the core requirement (Copilot Studio + Power Automate + Power Apps + Power BI), on a B2B contract that suits the Partita IVA. I did NOT submit because of one question: "What is your hands-on experience with Microsoft Copilot Studio? Describe the solutions you have built and your level of involvement in their design, development and deployment." You have the Copilot Studio Agent Academy and Power Automate certifications but I have no evidence you have SHIPPED a Copilot Studio solution, and I will not invent one. Tell me what you have actually built in it and I will finish the form. Also asks rate in PLN net/h + VAT (your 30-40k converts to roughly 75-100 PLN/h) and banking/fintech exposure.',
  'https://job-boards.greenhouse.io/capco/jobs/8002464'],
 ['Pigment','AI Deployment Strategist','London, hybrid','HELD','12 Aug 2026','-',
  'YOUR CALL. Best title match found in the whole sweep - "AI Deployment Strategist", 2-5 years, and they explicitly name Forward Deployed Engineer and Solutions Engineer as the backgrounds they want. Building and deploying custom AI agents for customers. But two hard fails stack: they list "Engineering or computer science degree from top tier institution" as REQUIRED, and London means UK sponsorship. You said chance the UK ones, so say the word and I will send it.',
  'https://jobs.lever.co/pigment/6cbc9f93-5034-4838-84b0-13fbcd742163'],
 ['Novosafe','GTM Automation & Growth Engineer','Berlin / EU remote','CLOSED','12 Aug 2026','-',
  'Search metadata said "updated 12 August 2026". The live page says archived 6 months ago. Would have been the closest match in the sweep - the posting named n8n by name. Letter is already written if they repost.',
  'https://join.com/companies/novosafede/15538213-gtm-automation-and-growth-engineer'],
 ['Arize AI','Forward Deployed AI Engineer, EMEA','Remote (EMEA)','CLOSED','12 Aug 2026','-',
  'Greenhouse redirected to the board index with ?error=true, the dead-posting signature. Checked the two EMEA roles still live on their board and dropped both: AI Solutions Manager EMEA wants 3+ yrs SaaS Customer Success and asks "Do you have 5 years of SaaS sales experience?", AI Sales Engineer EMEA is the same shape, and both require being based in the London area.',
  'https://job-boards.greenhouse.io/arizeai/jobs/5781408004'],
 ['Qdrant / Space Inch / Quadrivia / Action1 / Zeotap','Solutions Eng EMEA / AI Eng EU / AI Automation Eng / Applied AI Eng / Implementation Eng','EMEA, EU, ES-UK, Worldwide, London','CLOSED','12 Aug 2026','-',
  'All five dead on the live page. Qdrant "Job not found" and the rest of their board is Senior or US-only. Space Inch and Action1 both "This job is no longer available". Quadrivia "Job not found". Zeotap Lever 404. NOTE FOR NEXT TIME: three of these were verified only through RemoteRocketship mirrors rather than the ATS page itself, and all three were stale. Aggregator mirrors are not verification. One useful find: Action1 says on its board that it hires across the whole EU regardless of the country shown on a listing, and it has a live Backend Developer (Node.js/TS) role in Poland if you want a non-AI TypeScript option.',
  'https://apply.workable.com/action1/'],
]

ws.insert_rows(5, amount=len(NEW))
r = 5
for row in NEW:
    for i, v in enumerate(row, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.alignment = Alignment(vertical='top', wrap_text=(i in (2, 3, 6, 7)))
        c.border = BOX
        c.fill = FILL.get(row[3], GREY)
        if i == 4:
            c.font = Font(bold=True)
        if i == 8 and v:
            c.hyperlink = v
            c.value = 'open'
            c.font = Font(color='0563C1', underline='single')
    ws.row_dimensions[r].height = None
    r += 1

ws['A1'] = 'Applications sent - live log, last updated 12 August 2026 (evening: deep-dive batch added)'
ws['A1'].font = Font(bold=True, size=13)
wb.save('Role shortlist - 8 Aug 2026.xlsx')
print('inserted', len(NEW))
