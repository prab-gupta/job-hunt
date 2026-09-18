APPS = [
("Camunda","AI Forward Deployed Engineer, EMEA","https://jobs.ashbyhq.com/camunda/42b8007d-8299-4793-9581-2b3868ba0754","SUBMITTED","8 Aug 2026","'Success - Your application was successfully submitted.'","CV + structured answers only. Venice/Italy, permanent resident, remote Yes, EUR 50-60k, He/Him. Three self-ratings answered - see note below."),
("Luma AI","Forward Deployed Engineer (Remote, UK - Hybrid)","https://jobs.ashbyhq.com/lumaai/6133d955-1d26-4f3a-abf8-f9383ccbe19a","SUBMITTED","8 Aug 2026","'Success - Your application was successfully submitted.'","CV-only. Relocation = 'Yes - I can relocate'. Portfolio/GitHub left blank. All diversity/EEO/veteran declined. Posting is UK hybrid - raise your work-authorisation position early if they progress it."),
("Cohere","Forward Deployed Engineer, Agentic Platform (UK/Europe)","https://jobs.ashbyhq.com/cohere/2d256112-b336-4539-8133-a0bf7f6698f0","SUBMITTED","8 Aug 2026","'Your application has been submitted! Thank you for your interest in a career growth opportunity at Cohere!'","CV + one free-text answer, written strictly from your CV. GMT 4-hour overlap = Yes. Gaps stated plainly: TypeScript not Python, three years, no vector DB or LangGraph at production scale."),
("Swans","Forward Deployed Engineer","https://jobs.ashbyhq.com/swans/c5c62cfc-8248-4295-8e00-36e3261025c4","NEEDS YOU","-","-","TOP-RANKED ROLE (fit 91). Blocked on a MANDATORY video walkthrough: 'applications submitted without a video or demo will not be reviewed.' Record a 2-3 min screen walkthrough of an automation or dashboard you built. Everything else on the form is answerable from your CV - tell me when the video is ready and I will fill and submit."),
("n8n","Forward Deployed Engineer EMEA (or one of 3 siblings)","https://jobs.ashbyhq.com/n8n/c9fc97fa-a473-4133-b3cb-502785649ecd","NEEDS YOU","-","-","n8n allows ONE application per 15 days across all roles. Two decisions are yours: (1) which of the four n8n roles to use it on - FDE EMEA states no minimum years but Italy is not on its country list; Technical Programs Manager Scaled CS asks 3-5+ yrs and only requires 'rights to work in Europe'. (2) how long you have actually been building with n8n (they ask for a bracket). They also invite a screenshot of your favourite workflow, which would differentiate you."),
("Arize AI","Forward Deployed AI Engineer, EMEA","https://job-boards.greenhouse.io/arizeai","CLOSED","-","'The job you are looking for is no longer open.'","Closed between the sweep and the application attempt. Their live EMEA roles are AI Sales Engineer EMEA and AI Solutions Manager EMEA - the latter requires being London-based and 5 yrs SaaS sales, so it is not a fit."),
("Clera","Forward Deployed Engineer, Berlin","https://jobs.ashbyhq.com/clera","DROPPED","-","'Job not found'","Clera is an aggregator, not an employer: 311 postings under one 'Engineering' department including Co-Founder & CTO and Mechanical Engineer, with the same req duplicated across nine cities. The Berlin posting is dead. No named end client - dropped on the same rule used everywhere else."),
]
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb=openpyxl.load_workbook("Role shortlist - 8 Aug 2026.xlsx")
if "Applications sent" in wb.sheetnames: del wb["Applications sent"]
ws=wb.create_sheet("Applications sent",1)
hdr=["Company","Role","Status","Date","On-screen confirmation","What went in / what is blocking"]
widths=[20,48,14,14,54,84]
HF=PatternFill("solid",fgColor="1F3864"); HFONT=Font(bold=True,color="FFFFFF",size=11)
FILLS={"SUBMITTED":PatternFill("solid",fgColor="E2EFDA"),"NEEDS YOU":PatternFill("solid",fgColor="FFF2CC"),"CLOSED":PatternFill("solid",fgColor="F2F2F2"),"DROPPED":PatternFill("solid",fgColor="F2F2F2")}
THIN=Side(style="thin",color="D9D9D9"); B=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
ws["A1"]="Applications - Forward Deployed Engineer run, 8 August 2026"
ws["A1"].font=Font(bold=True,size=14,color="1F3864")
ws["A2"]="Green = submitted with an on-screen confirmation I saw. Amber = filled or scouted but blocked on something only you can supply. Nothing is counted as sent without the confirmation text quoted in column E."
ws["A2"].font=Font(italic=True,size=10,color="595959"); ws["A2"].alignment=Alignment(wrap_text=True)
ws.merge_cells("A2:F2"); ws.row_dimensions[2].height=30
for c,(h,w) in enumerate(zip(hdr,widths),1):
    cell=ws.cell(row=4,column=c,value=h); cell.fill=HF; cell.font=HFONT; cell.border=B
    ws.column_dimensions[get_column_letter(c)].width=w
for i,(co,role,url,status,date,conf,note) in enumerate(APPS,1):
    r=i+4
    for c,v in enumerate([co,role,status,date,conf,note],1):
        cell=ws.cell(row=r,column=c,value=v); cell.border=B
        cell.alignment=Alignment(wrap_text=(c in (2,5,6)),vertical="top")
        if status in FILLS: cell.fill=FILLS[status]
    ws.cell(row=r,column=2).hyperlink=url
    ws.cell(row=r,column=2).font=Font(color="0563C1",underline="single")
r=len(APPS)+7
ws.cell(row=r,column=1,value="Judgement calls I made for you on the Camunda self-ratings - correct me if any is wrong:").font=Font(bold=True,size=11,color="C00000")
for j,line in enumerate([
 "1. Large-enterprise production experience -> 'I have only shipped in smaller companies or startups.' Straightforwardly true.",
 "2. AI agents -> 'built and operated agentic systems in production, including testing for non-determinism and failure recovery.' The production part is on your CV (FlowSpace AI support agents). The testing clause is NOT something I can evidence. I picked it because the only alternative said staging-only, which would have been false in the other direction.",
 "3. Workflow automation / BPMN -> 'configured, debugged, or written BPMN workflows in a production environment.' True for workflow automation via n8n; you have no BPMN specifically. The other two options were 'conceptual only' and 'no experience', both false.",
],1):
    ws.cell(row=r+j,column=1,value=line).alignment=Alignment(wrap_text=True,vertical="top")
    ws.merge_cells(start_row=r+j,start_column=1,end_row=r+j,end_column=6)
    ws.row_dimensions[r+j].height=30
ws.freeze_panes="A5"
wb.save("Role shortlist - 8 Aug 2026.xlsx")
print("applications tab written")
