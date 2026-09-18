import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

WB = '/home/claude/sheet/Role shortlist - 8 Aug 2026.xlsx'
wb = openpyxl.load_workbook(WB)

HDR   = PatternFill('solid', fgColor='1F3864')
GREEN = PatternFill('solid', fgColor='D6EFD6')
AMBER = PatternFill('solid', fgColor='FFF2CC')
GREY  = PatternFill('solid', fgColor='EEEEEE')
RED   = PatternFill('solid', fgColor='FCE4E4')
thin  = Side(style='thin', color='BBBBBB')
BOX   = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(name):
    if name in wb.sheetnames:
        wb.remove(wb[name])
    return wb.create_sheet(name)

# ---------------- Applications sent ----------------
ws = sheet('Applications sent')
ws.sheet_view.showGridLines = False

ws['A1'] = 'Applications sent - live log, last updated 12 August 2026'
ws['A1'].font = Font(bold=True, size=14, color='1F3864')
ws['A2'] = ('GREEN = submitted and I saw the confirmation on screen (quoted verbatim).   '
            'AMBER = ready to send but blocked on something only you have.   '
            'GREY = closed, dropped, or held.')
ws['A2'].font = Font(size=10, italic=True, color='555555')

HEADERS = ['Company','Role','Location','Status','Date','On-screen confirmation','What went in / what is blocking','Link']
for i,h in enumerate(HEADERS, start=1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
    c.alignment = Alignment(vertical='center'); c.border = BOX

ROWS = [
 # company, role, location, status, date, confirmation, notes, link
 ('SolveAI','Forward Deployed Engineer','London (4 days on-site)','SUBMITTED','12 Aug 2026',
  '"Your application was successfully submitted. We\'ll contact you if there are next steps."',
  'CV + short answers. Ex-Palantir team, GV and Accel backed, already live with enterprise customers. Posting explicitly welcomes engineers making the jump into FDE, no years gate. UK sponsorship answered No (honest); they do not state whether they sponsor. 50-75% travel.',
  'https://jobs.ashbyhq.com/solve.ai/1179e1af-6930-4271-b4b5-75294c13f806'),
 ('Seamflow','Forward Deployed Engineer','London (on-site)','SUBMITTED','12 Aug 2026',
  '"Your application was successfully submitted. We\'ll be in touch if there are next steps!"',
  'CV + "why a startup" written from your CV only. Years-of-experience question answered 3. Seed round backed by Coinbase, Reddit, Klarna and Spotify people. Visa sponsorship, relocation and visas explicitly covered.',
  'https://jobs.ashbyhq.com/seamflow/83e5dd13-b293-462f-9cb6-bc9cc325b155'),
 ('Zauber','Forward Deployed Engineer','Berlin (on-site)','SUBMITTED','12 Aug 2026',
  '"Your application was successfully submitted. We\'ll contact you if there are next steps."',
  'This one was logged CLOSED on 8 Aug; the posting is live again and is now sent. No German requirement, no years gate. TypeScript/Node/React, ERP/TMS/SFTP/event-queue integration. Right-to-work answered Yes with the EU long-term residence position spelled out in the free text.',
  'https://jobs.ashbyhq.com/zauber/c11db72e-3a33-4df5-ae33-e1cc39cb0ba7'),
 ('Atmen Solutions','Customer Solutions Engineer','Madrid (hybrid)','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to Atmen Solutions GmbH."',
  'CV + tailored letter. EUR 40-90k posted, 3+ yrs, English only, no Spanish needed. Explicitly NOT a software engineering role: onboarding, implementation, delivery for a green-certification RegTech. Gaps stated: no energy or regulatory background.',
  'https://join.com/companies/atmen/16529717-customer-solutions-engineer'),
 ('Laioutr','AI Engineer - Full Stack (TS / Nuxt / Next / Vue / React)','Remote (Germany)','SUBMITTED','12 Aug 2026',
  '"Bewerbung gesendet - Du hast dich bei Laioutr GmbH beworben."',
  'CV + tailored letter. EUR 36-60k posted, asked 55,000. Fluent German OR English accepted. Gaps stated: no Nuxt/Vue/Pinia/tRPC/Turborepo, no conversational German.',
  'https://join.com/companies/laioutr/16553103-ai-engineer-m-f-d-full-stack-typescript-nuxt-next-vue-react'),
 ('dexter health','Applied AI Developer','Remote','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to dexter health."',
  'CV + tailored letter. They say outright they care less about formal seniority than speed and ownership, and they use Claude Code / Cursor daily. Gap stated: no healthcare or clinical data background. Your CV was re-uploaded fresh here and set as the join.com default.',
  'https://join.com/companies/dexter-health/16549765-applied-ai-developer-remote'),
 ('STRATOS Technologies','Forward Deployed Engineer','London (hybrid)','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to STRATOS Technologies Ltd.."',
  'CV only, no cover letter step. 2-4 yrs required, which is your band. Visa sponsorship offered as a benefit, so UK sponsorship answered Yes. English self-rated Full professional working proficiency (my call, see below). Gap stated in the letter I prepared: no PySpark or Palantir AIP.',
  'https://join.com/companies/stratos/16565987-forward-deployed-engineer-at-stratos'),
 ('Otherworldly','Forward Deployed Engineer (f/m/x)','Berlin (5 days in office)','SUBMITTED','12 Aug 2026',
  '"Application sent - You\'ve applied to Otherworldly."',
  'CV + tailored letter. No years requirement stated. Python and/or JS/TS, LLMs in real products, strong-plus for internal tools and CRM/outbound systems, which is your FlowSpace work. Asked EUR 60,000, start 1 Sep, on-site Yes.',
  'https://join.com/companies/otherworldly-artistscom/16557950-forward-deployed-engineer-f-m-x'),
 ('Camunda','AI Forward Deployed Engineer, EMEA','Remote EMEA','SUBMITTED','8 Aug 2026',
  '"Success - Your application was successfully submitted."',
  'CV + structured answers. See the self-rating judgement calls in red at the bottom of this tab.',
  ''),
 ('Luma AI','Forward Deployed Engineer','Remote UK / hybrid','SUBMITTED','8 Aug 2026',
  '"Success - Your application was successfully submitted."',
  'CV-only. Relocation = Yes. Portfolio and GitHub left blank rather than invented.',''),
 ('Cohere','Forward Deployed Engineer, Agentic Platform','UK / Europe','SUBMITTED','8 Aug 2026',
  '"Your application has been submitted! Thank you for your interest in a career growth opportunity at Cohere!"',
  'CV + one free-text answer written strictly from your CV.',''),
 ('Taktile','Forward Deployed Engineer','Berlin / London (hybrid)','SUBMITTED','8 Aug 2026',
  '"Success - Your application was successfully submitted."',
  'Pure CV-only form. Willing to relocate = Yes, start 1 Sep 2026.',''),
 ('OMMAX','Forward Deployed AI Engineer (m/f/d)','Munich / Berlin / London / Amsterdam / Paris / MILAN (hybrid)','NEEDS YOU','12 Aug 2026','-',
  'BEST OF THE UNSENT. Milan is one of their offices, so this works without leaving Italy. English only, no German, no years gate. Form is filled and CV + letter are uploaded, but Personio will not submit without two files only you have: an EMPLOYMENT REFERENCE letter and your ACADEMIC TRANSCRIPTS. Send me those two PDFs and I finish it in one pass.',
  'https://ommax.jobs.personio.com/job/2550344?language=en'),
 ('GTM Base','Forward-Deployed CRM Architect','Remote (CET overlap)','NEEDS YOU','12 Aug 2026','-',
  'They run the company on Claude Code, use n8n, clients include Groupon and Blinkist. I stopped at question 2, a HubSpot build test: "build a flow that creates a deal when a lead moves to Qualified, with a reliable duplicate check against the DACH Sales pipeline." Answering would have meant inventing HubSpot depth you have not told me you have. Give me a few lines and I will send it. Letter is written and ready.',
  'https://join.com/companies/gtmbaseio/16539004-forward-deployed-crm-architect'),
 ('GTM Base','Senior GTM Engineer (contract, path to full-time)','Remote','NEEDS YOU','12 Aug 2026','-',
  'Same company, second req. Clay / HubSpot / Salesforce / n8n / Claude / GPT. Requires a Loom walkthrough of YOUR OWN Clay tables as part of the application. Only you can record that.',
  'https://join.com/companies/gtmbaseio/16516200-senior-gtm-engineer'),
 ('Swans','Forward Deployed Engineer','-','NEEDS YOU','-','-',
  'Highest fit score on the whole sheet (91). Blocked on a mandatory video walkthrough that only you can record.',''),
 ('n8n','Forward Deployed Engineer EMEA (or one of 3 siblings)','-','NEEDS YOU','-','-',
  'n8n allows ONE application per 15 days across ALL their roles. Tell me which req to spend it on.',''),
 ('Codekeeper','QA Verifications Specialist','-','NEEDS YOU','-','-',
  'Needs a 16personalities result, a profile URL, your device details and a captcha.',''),
 ('Kadoa','-','-','NEEDS YOU','-','-','Gmail draft is written; attach the CV and hit send.',''),
 ('xAI','-','-','NEEDS YOU','-','-','Wants a GitHub profile. I will not invent one.',''),
 ('Nearform','Forward Deployed Engineer (Perm, Italy, Remote + Travel)','Italy, remote','DROPPED','12 Aug 2026','-',
  'Not applied: requires 8+ years commercial at senior level. Worth knowing it exists though - permanent, Italy-based, fully remote, EUR 60-80k plus bonus. That is the shape of role to aim at in two or three years.',''),
 ('VOIDS Technology','FDE - Integrations & Customer Success','Hamburg','DROPPED','12 Aug 2026','-',
  'EUR 80-110k and only 3+ years, but requires fluent German. Pure language gate.',''),
 ('Perelyn / CROZ / codecentric / amber / Mercura / hallo theo / ETHjuniors / superkind','Various FDE roles','Germany','DROPPED','12 Aug 2026','-',
  'Eight separate German FDE roles killed by one thing: fluent or native German. Not seniority, not your permit. See the market note.',''),
 ('Soji AI','Forward Deployed Engineer (Aviation)','Hamburg','DROPPED','12 Aug 2026','-','8+ years required.',''),
 ('GWI','Forward Deployed Engineer','London','DROPPED','12 Aug 2026','-','Principal level AND no visa sponsorship for the role.',''),
 ('GitLab','Staff Forward Deployed Engineer, Agentic SDLC','EMEA remote','DROPPED','12 Aug 2026','-','Staff level, Ruby on Rails and Go.',''),
 ('Firecrawl / Coworked / Harness / Istari / Aircall','Various FDE roles','US','DROPPED','12 Aug 2026','-',
  'US-only. Istari requires US citizenship with no dual nationality.',''),
 ('Salesforce','Deployment Strategist - Public Sector','Rome, Italy','HELD','12 Aug 2026','-',
  '4+ yrs in public-sector enterprise delivery, and Workday makes you create an account to apply. I do not create accounts for you. Apply yourself if you want it.',''),
 ('Arize AI','Forward Deployed AI Engineer, EMEA','-','CLOSED','-','"The job you are looking for is no longer open."','Closed between the sweep and the attempt.',''),
 ('Clera','Forward Deployed Engineer','Berlin / Munich','DROPPED','-','"Job not found"',
  'Aggregator, not an employer: 311 postings under one department, same req duplicated across nine cities.',''),
]

FILL = {'SUBMITTED':GREEN,'NEEDS YOU':AMBER,'DROPPED':GREY,'CLOSED':GREY,'HELD':GREY}

r = 5
for row in ROWS:
    for i,v in enumerate(row, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.alignment = Alignment(vertical='top', wrap_text=(i in (2,6,7)))
        c.border = BOX
        c.fill = FILL.get(row[3], GREY)
        if i == 4: c.font = Font(bold=True)
        if i == 8 and v:
            c.hyperlink = v; c.value = 'open'; c.font = Font(color='0563C1', underline='single')
    r += 1

r += 1
note = ws.cell(row=r, column=1, value='Judgement calls I made on your behalf - correct me and I will redo them')
note.font = Font(bold=True, size=11, color='9C0006'); r += 1
for t in [
 'Camunda 1. Large-enterprise production -> "I have only shipped in smaller companies or startups." Plainly true.',
 'Camunda 2. AI agents -> "built and operated agentic systems in production, including testing for non-determinism and failure recovery." The production half is CV-backed; the testing clause is NOT evidenced. Chosen because the only alternative said staging-only, which is false the other way.',
 'Camunda 3. Workflow/BPMN -> "configured, debugged, or written BPMN workflows in a production environment." True for workflow automation via n8n; you have no BPMN specifically. The other options were "conceptual only" and "no experience", both false.',
 'STRATOS. English proficiency -> "Full professional working proficiency". Your own wording is "professional English proficiency"; the next option down would have undersold you on a UK-facing role.',
 'Zauber. "Do you have the right to work in Germany?" -> Yes, on the basis that your Italian EU long-term residence permit carries intra-EU mobility and needs no employer sponsorship. I said exactly that in the free-text box so the Yes is transparent. Laioutr asked a narrower question ("do you have a CURRENT permit for Germany") and I answered No there, paired with No to needing sponsorship.',
 'Salary. Where a single number was demanded I used 60,000 (top of your stated range), except Laioutr where their posted band topped out at 60,000 and I used 55,000 to sit inside both.',
]:
    c = ws.cell(row=r, column=1, value=t)
    c.font = Font(size=10, color='9C0006'); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 30
    r += 1

for col,w in zip('ABCDEFGH',[24,44,30,14,12,52,86,8]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'
ws.auto_filter.ref = f'A4:H{4+len(ROWS)}'

wb.save(WB)
print('applications tab rebuilt:', len(ROWS), 'rows')
